"""Parse the ITU E.212 MNC list annexed to OB No. 1280 (position 15 Nov 2023).

Usage: python3 -I parse_baseline.py <docx> <out_json>
"""
import json, re, sys
import docx

CODE_RE = re.compile(r"^\s*(\d{3})\s+(\d{2,3})\s*$")


def clean(s):
    return re.sub(r"\s+", " ", s.replace("\xa0", " ")).strip()


def main(path, out):
    d = docx.Document(path)
    rows, problems = [], []

    # Table 1: geographic MCCs
    t = d.tables[1]
    for i, r in enumerate(t.rows[1:], start=1):
        c = [clean(x.text) for x in r.cells]
        country, network, code = c[0], c[1], c[2]
        if not network and not code:
            continue  # country header row
        m = CODE_RE.match(code)
        if not m:
            problems.append({"table": 1, "row": i, "cells": c})
            continue
        rows.append({
            "mcc": m.group(1), "mnc": m.group(2), "country": country,
            "network": network, "scope": "geographic",
        })

    # Table 2: MCC 901 international shared code
    for i, r in enumerate(d.tables[2].rows[1:], start=1):
        c = [clean(x.text) for x in r.cells]
        m = CODE_RE.match(c[1])
        if not m:
            problems.append({"table": 2, "row": i, "cells": c})
            continue
        rows.append({
            "mcc": m.group(1), "mnc": m.group(2), "country": "International (shared code)",
            "network": c[0], "scope": "international",
        })

    # Table 3: MCC 902 ROIO/SDO shared code
    for i, r in enumerate(d.tables[3].rows[1:], start=1):
        c = [clean(x.text) for x in r.cells]
        if not c[2]:
            continue
        m = CODE_RE.match(c[2])
        if not m:
            problems.append({"table": 3, "row": i, "cells": c})
            continue
        rows.append({
            "mcc": m.group(1), "mnc": m.group(2), "country": "ROIO/SDO-specified (shared code)",
            "network": c[1], "scope": "roio_sdo",
        })

    # Table 4: extra-territorial use
    extra = []
    for i, r in enumerate(d.tables[4].rows[1:], start=1):
        c = [clean(x.text) for x in r.cells]
        if not any(c):
            continue
        extra.append({"country_role": c[0], "code": c[1], "msin_range": c[2], "operator": c[3]})

    json.dump({
        "source": "ITU-T E.212 MNC list, annex to Operational Bulletin No. 1280",
        "position_date": "2023-11-15",
        "codes": rows, "extra_territorial": extra, "parse_problems": problems,
    }, open(out, "w"), ensure_ascii=False, indent=1)

    from collections import Counter
    print("codes:", len(rows), "| by scope:", dict(Counter(r["scope"] for r in rows)))
    print("extra-territorial rows:", len(extra))
    print("parse problems:", len(problems))
    for p in problems[:20]:
        print("  ", p)
    dup = Counter((r["mcc"], r["mnc"]) for r in rows)
    dups = [k for k, v in dup.items() if v > 1]
    print("duplicate codes:", len(dups), dups[:20])
    print("countries:", len({r['country'] for r in rows if r['scope'] == 'geographic'}))
    print("distinct MCCs:", len({r['mcc'] for r in rows}))


if __name__ == "__main__":
    main(*sys.argv[1:3])
